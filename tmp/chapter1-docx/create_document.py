from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from uuid import uuid4
from lxml import etree as ET
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/Users/initia-army/Documents/juristic-platform')
OUT=ROOT/'output/documents/บทที่1_ที่มา_วัตถุประสงค์_ขอบเขตโครงงาน.docx'
FONT='TH Sarabun New'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.54);sec.bottom_margin=Cm(2.54)
sec.left_margin=Cm(3);sec.right_margin=Cm(2.54)
sec.footer_distance=Cm(1.25)

for sty in doc.styles:
 if sty.type in (1,2):
  sty.font.name=FONT
  sty.font.color.rgb=RGBColor(0,0,0)
  rp=sty.element.get_or_add_rPr()
  for child in list(rp):
   if child.tag in [qn('w:'+x) for x in ('szCs','bCs','i','iCs','spacing','kern')]:rp.remove(child)
  pp=sty.element.find(qn('w:pPr'))
  if pp is not None:
   for child in list(pp):
    if child.tag in [qn('w:'+x) for x in ('pBdr','numPr')]:pp.remove(child)
  rf=rp.find(qn('w:rFonts'))
  if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
  for a in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+a),FONT)
  for a in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
   rf.attrib.pop(qn('w:'+a),None)
  lang=OxmlElement('w:lang');lang.set(qn('w:val'),'th-TH');lang.set(qn('w:eastAsia'),'th-TH');lang.set(qn('w:bidi'),'th-TH');rp.append(lang)

for name,size,bold in [('Normal',16,False),('Title',20,True),('Subtitle',18,True),('Heading 1',18,True),('Heading 2',16,True)]:
 sty=doc.styles[name];sty.font.size=Pt(size);sty.font.bold=bold;sty.font.italic=False
 rp=sty.element.get_or_add_rPr()
 sz=OxmlElement('w:szCs');sz.set(qn('w:val'),str(size*2));rp.append(sz)
 if bold:
  bc=OxmlElement('w:bCs');rp.append(bc)
 pf=sty.paragraph_format
 pf.line_spacing=Pt(19 if name=='Normal' else 21)
 pf.space_after=Pt(6)
 if name.startswith('Heading'):
  pf.space_before=Pt(10);pf.space_after=Pt(5);pf.keep_with_next=True
 if name in ('Title','Subtitle'):
  pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
  pf.space_before=Pt(0);pf.space_after=Pt(8);pf.keep_with_next=True
normal=doc.styles['Normal'].paragraph_format
normal.widow_control=True

# The Thai run language and all four font slots keep complex-script text consistent in Word.
def para(text, indent=True):
 p=doc.add_paragraph(text)
 p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if indent:p.paragraph_format.first_line_indent=Cm(1.25)
 return p

def item(n,text):
 p=doc.add_paragraph(f'{n})  {text}')
 p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 p.paragraph_format.left_indent=Cm(.7)
 p.paragraph_format.first_line_indent=Cm(-.7)
 p.paragraph_format.space_after=Pt(4)
 p.paragraph_format.keep_together=True
 return p

def heading(text,level=1):
 p=doc.add_paragraph(text,style=f'Heading {level}')
 if text in ('วัตถุประสงค์ของโครงงาน','ขอบเขตด้านข้อมูลและการทำงานของระบบ','ขอบเขตด้านเทคโนโลยีและการทดสอบ'):
  p.paragraph_format.page_break_before=True
 return p

doc.add_paragraph('บทที่ 1\nบทนำ',style='Title')
doc.add_paragraph('การพัฒนาเว็บแอปพลิเคชันสำหรับจัดการค่าสาธารณูปโภค\nและค่าส่วนกลางของอพาร์ตเมนต์\nพร้อมระบบแจ้งเตือนผ่าน LINE Official Account',style='Subtitle')

heading('ที่มาและความสำคัญของโครงงาน')
para('การจัดเก็บค่าน้ำ ค่าไฟฟ้า และค่าส่วนกลางเป็นงานที่ผู้ดูแลอพาร์ตเมนต์ต้องดำเนินการเป็นประจำทุกเดือน โดยเริ่มจากการบันทึกเลขมิเตอร์ของแต่ละห้อง การคำนวณค่าใช้จ่าย การออกใบแจ้งหนี้ และการติดตามการชำระเงิน ข้อมูลในแต่ละขั้นตอนมีความสัมพันธ์กัน หากบันทึกข้อมูลไม่ครบถ้วนหรือใช้ข้อมูลคนละช่วงเวลา อาจทำให้ยอดเรียกเก็บคลาดเคลื่อนและต้องใช้เวลาในการตรวจสอบแก้ไข')
para('การดำเนินงานด้วยเอกสารหรือตารางคำนวณที่แยกจากกันอาจทำให้ต้องกรอกข้อมูลซ้ำ ค้นหาประวัติการใช้น้ำและไฟฟ้าได้ไม่สะดวก และติดตามได้ยากว่าห้องใดชำระเงินแล้วหรือยังค้างชำระ นอกจากนี้ การแจ้งยอดและแจ้งเตือนผู้เช่าเป็นรายห้องยังเพิ่มภาระงานของผู้ดูแล โดยเฉพาะเมื่อต้องตรวจสอบข้อมูลย้อนหลังหรือจัดทำสรุปยอดประจำเดือน จึงควรมีระบบที่รวบรวมข้อมูลและเชื่อมโยงขั้นตอนดังกล่าวไว้ในแหล่งเดียวกัน')
para('ผู้จัดทำจึงมีแนวคิดพัฒนาเว็บแอปพลิเคชันสำหรับจัดการค่าสาธารณูปโภคและค่าส่วนกลางของอพาร์ตเมนต์ โดยมุ่งให้ผู้ดูแลและเจ้าหน้าที่ใช้งานผ่านเว็บเบราว์เซอร์บนคอมพิวเตอร์หรือโทรศัพท์มือถือ ระบบจะจัดเก็บข้อมูลห้องพัก ผู้เช่า และเลขมิเตอร์ คำนวณค่าน้ำและค่าไฟฟ้าจากจำนวนหน่วยที่ใช้ และคำนวณค่าส่วนกลางจากพื้นที่ห้องพักตามอัตราที่กำหนด ก่อนจัดทำใบแจ้งหนี้ให้ผู้ดูแลตรวจสอบและอนุมัติ รวมถึงบันทึกการรับชำระเงินสดและออกใบเสร็จรับเงิน')
para('โครงงานยังมีแนวทางเชื่อมต่อ LINE Official Account เพื่อส่งใบแจ้งหนี้ การแจ้งเตือนก่อนครบกำหนดและเมื่อค้างชำระ ข้อความยืนยันการชำระเงิน และประกาศถึงผู้เช่า ควบคู่กับการแสดงข้อมูลสรุปและประวัติย้อนหลัง เพื่อช่วยให้ผู้ดูแลติดตามการเรียกเก็บได้สะดวก ลดโอกาสผิดพลาดจากการคำนวณด้วยตนเอง และมีข้อมูลประกอบการตรวจสอบรายการรับชำระเงินอย่างเป็นระบบ')

heading('วัตถุประสงค์ของโครงงาน')
objectives=[
'เพื่อพัฒนาเว็บแอปพลิเคชันที่รวบรวมข้อมูลห้องพัก ผู้เช่า ค่าสาธารณูปโภค และค่าส่วนกลางไว้ในฐานข้อมูลเดียวกัน',
'เพื่อพัฒนาระบบบันทึกเลขมิเตอร์และคำนวณค่าน้ำ ค่าไฟฟ้า และค่าส่วนกลางตามอัตราที่กำหนด พร้อมจัดทำใบแจ้งหนี้รายเดือน',
'เพื่อพัฒนาระบบอนุมัติใบแจ้งหนี้ บันทึกการรับชำระเงินสด ออกใบเสร็จรับเงิน และตรวจสอบประวัติการชำระเงิน',
'เพื่อพัฒนาการแจ้งยอด แจ้งเตือนการชำระเงิน และส่งประกาศผ่าน LINE Official Account แก่ผู้เช่าที่ผูกบัญชีกับระบบ',
'เพื่อจัดทำหน้าสรุปข้อมูลและรายงานสำหรับติดตามยอดเรียกเก็บ ยอดรับชำระ ยอดค้างชำระ และประวัติการใช้น้ำและไฟฟ้า',
'เพื่อทดสอบความถูกต้องของการคำนวณ ลำดับการทำงาน และการกำหนดสิทธิ์ของผู้ใช้งานตามขอบเขตที่กำหนด'
]
for n,t in enumerate(objectives,1):item(n,t)

heading('ขอบเขตของโครงงาน')
para('โครงงานมุ่งพัฒนาระบบสำหรับอพาร์ตเมนต์หนึ่งแห่งและหนึ่งอาคาร ครอบคลุมค่าน้ำ ค่าไฟฟ้า และค่าส่วนกลางรายเดือน ตั้งแต่บันทึกเลขมิเตอร์จนถึงออกใบเสร็จรับเงิน โดยมีขอบเขตดังต่อไปนี้')

heading('ขอบเขตด้านผู้ใช้งาน',2)
item(1,'ผู้ดูแลระบบหรือนิติบุคคล สามารถจัดการข้อมูลห้องพักและผู้เช่า กำหนดอัตราค่าบริการ ตรวจสอบและอนุมัติใบแจ้งหนี้ บันทึกการรับชำระเงิน ออกใบเสร็จรับเงิน จัดทำประกาศ และดูข้อมูลสรุปกับรายงานทางการเงิน')
item(2,'เจ้าหน้าที่หรือช่าง สามารถดูรายการห้องพักและเลขมิเตอร์ครั้งก่อน บันทึกเลขมิเตอร์น้ำและไฟฟ้าประจำเดือน และตรวจสอบความครบถ้วนของการบันทึกตามสิทธิ์ที่ได้รับ โดยไม่มีสิทธิ์เข้าถึงข้อมูลทางการเงินหรืออนุมัติใบแจ้งหนี้')
item(3,'ผู้เช่าหรือลูกบ้าน รับใบแจ้งหนี้ การแจ้งเตือน ข้อความยืนยันการชำระเงิน และประกาศผ่าน LINE Official Account หลังผูกบัญชีกับห้องพัก โดยไม่ต้องมีบัญชีเว็บสำหรับเจ้าหน้าที่')

heading('ขอบเขตด้านข้อมูลและการทำงานของระบบ',2)
features=[
'จัดเก็บข้อมูลพื้นฐาน ได้แก่ หมายเลขห้อง ชั้น ประเภทห้อง พื้นที่ห้อง สถานะห้อง ผู้เช่าปัจจุบัน และข้อมูลติดต่อ รวมถึงข้อมูลสำหรับเชื่อมโยงผู้เช่ากับบัญชี LINE',
'จัดเก็บเลขมิเตอร์น้ำและไฟฟ้าแยกตามห้องและรอบเดือน พร้อมข้อมูลผู้บันทึก ชื่อช่างที่อ่านมิเตอร์ และวันเวลาที่บันทึก โดยตรวจสอบความครบถ้วนและความสมเหตุสมผลของเลขมิเตอร์ก่อนนำไปคำนวณ',
'คำนวณจำนวนหน่วยน้ำและไฟฟ้าจากผลต่างระหว่างเลขมิเตอร์ปัจจุบันกับเลขมิเตอร์ครั้งก่อน แล้วคูณด้วยอัตราค่าบริการต่อหน่วย ส่วนค่าส่วนกลางคำนวณจากพื้นที่ห้องพักคูณด้วยอัตราต่อตารางเมตรที่กำหนด',
'จัดทำใบแจ้งหนี้รายเดือนของแต่ละห้อง โดยแสดงรายละเอียดค่าน้ำ ค่าไฟฟ้า ค่าส่วนกลาง ยอดรวม และวันครบกำหนดชำระ พร้อมกระบวนการตรวจสอบและอนุมัติ และแสดงสถานะร่าง รอชำระ หรือชำระแล้ว รวมถึงระบุรายการที่เกินกำหนดชำระ',
'บันทึกการรับชำระเงินสดที่สำนักงานตามใบแจ้งหนี้ โดยจัดเก็บจำนวนเงิน วันเวลา และผู้รับเงิน พร้อมปรับสถานะการชำระเงินและออกใบเสร็จที่สามารถเรียกดูหรือพิมพ์ได้',
'พัฒนาการเชื่อมต่อ LINE Official Account สำหรับส่งใบแจ้งหนี้ แจ้งเตือนก่อนครบกำหนด แจ้งเตือนยอดค้างชำระ ส่งข้อความยืนยันการรับเงิน และส่งประกาศ โดยอาศัยข้อมูลการผูกบัญชีผู้เช่ากับห้องพัก',
'แสดงหน้าสรุปจำนวนห้อง ยอดเรียกเก็บ ยอดรับชำระ และยอดค้างชำระ พร้อมกราฟหรือประวัติการใช้น้ำและไฟฟ้า ตลอดจนรายงานสำหรับค้นหาห้องที่ยังไม่ชำระเงิน',
'กำหนดสิทธิ์การเข้าถึงตามบทบาทของผู้ใช้งาน และจัดเก็บประวัติการดำเนินงานที่สำคัญ เช่น การบันทึกมิเตอร์ การอนุมัติใบแจ้งหนี้ และการบันทึกรับชำระเงิน เพื่อใช้ตรวจสอบย้อนหลัง'
]
for n,t in enumerate(features,1):item(n,t)

heading('ขอบเขตด้านเทคโนโลยีและการทดสอบ',2)
para('พัฒนาระบบในรูปแบบเว็บแอปพลิเคชันที่รองรับการแสดงผลบนคอมพิวเตอร์และโทรศัพท์มือถือ โดยใช้ Node.js และ Express สำหรับการทำงานฝั่งเซิร์ฟเวอร์ ใช้ EJS สำหรับสร้างหน้าเว็บ และใช้ฐานข้อมูล PostgreSQL สำหรับจัดเก็บข้อมูล ส่วนการแจ้งเตือนผู้เช่ากำหนดให้เชื่อมต่อผ่าน LINE Official Account',False)
para('การทดสอบใช้ข้อมูลตัวอย่างเพื่อตรวจสอบการคำนวณค่าใช้จ่าย การสร้างและอนุมัติใบแจ้งหนี้ การรับชำระเงินและออกใบเสร็จ การเรียกดูข้อมูลย้อนหลัง และการจำกัดสิทธิ์ของผู้ใช้งาน สำหรับส่วนเชื่อมต่อ LINE จะทดสอบการผูกบัญชีและการส่งข้อความตามเหตุการณ์ที่กำหนดเมื่อพัฒนาการเชื่อมต่อแล้ว',False)

heading('ข้อจำกัดและงานที่ไม่รวมในโครงงาน',2)
item(1,'ไม่ครอบคลุมการคิดค่าเช่าห้อง การจัดทำสัญญาเช่า และการบริหารหลายอพาร์ตเมนต์หรือหลายอาคารภายในระบบเดียวกัน')
item(2,'รองรับการบันทึกรับชำระเงินสดที่สำนักงานเท่านั้น ไม่รวมการชำระเงินออนไลน์ การเชื่อมต่อธนาคารหรือผู้ให้บริการรับชำระเงิน การชำระผ่าน QR Code และการตรวจสอบสลิปโอนเงิน')
item(3,'เจ้าหน้าที่เป็นผู้กรอกเลขมิเตอร์ด้วยตนเอง ไม่รวมอุปกรณ์อ่านมิเตอร์อัตโนมัติ การเชื่อมต่ออุปกรณ์ IoT หรือการวิเคราะห์ข้อมูลด้วยปัญญาประดิษฐ์')
item(4,'ไม่พัฒนาแอปพลิเคชันสำหรับติดตั้งบนโทรศัพท์มือถือโดยเฉพาะ และการส่งข้อความถึงผู้เช่าขึ้นอยู่กับการผูกบัญชี LINE การเชื่อมต่ออินเทอร์เน็ต และเงื่อนไขการให้บริการของ LINE Official Account')

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(14)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='บทที่ 1 บทนำ การพัฒนาเว็บแอปพลิเคชันจัดการค่าสาธารณูปโภคและค่าส่วนกลางของอพาร์ตเมนต์'
doc.core_properties.subject='ที่มาและความสำคัญ วัตถุประสงค์ และขอบเขตโครงงาน'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.save(OUT)

# Embed complete regular and bold font files so the requested typeface travels with the document.
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
with ZipFile(OUT) as z:parts={n:z.read(n) for n in z.namelist()}
fonts=ET.fromstring(parts['word/fontTable.xml'])
font=next((e for e in fonts if e.get('{'+W+'}name')==FONT),None)
if font is None:font=ET.SubElement(fonts,'{'+W+'}font',{'{'+W+'}name':FONT})
rels=ET.Element('{'+REL+'}Relationships',nsmap={None:REL})
for i,(kind,filename) in enumerate([('embedRegular','THSarabunNew.ttf'),('embedBold','THSarabunNew-Bold.ttf')],1):
 key=uuid4();data=bytearray((ROOT/'tmp/chapter1-docx/fonts'/filename).read_bytes())
 xor=key.bytes[::-1]
 for j in range(32):data[j]^=xor[j%16]
 target=f'fonts/font{i}.odttf';parts['word/'+target]=bytes(data)
 ET.SubElement(font,'{'+W+'}'+kind,{'{'+R+'}id':f'rId{i}','{'+W+'}fontKey':'{'+str(key).upper()+'}','{'+W+'}subsetted':'false'})
 ET.SubElement(rels,'{'+REL+'}Relationship',Id=f'rId{i}',Type=R+'/font',Target=target)
parts['word/fontTable.xml']=ET.tostring(fonts,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/_rels/fontTable.xml.rels']=ET.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=ET.fromstring(parts['[Content_Types].xml'])
ET.SubElement(ct,'{'+CT+'}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
parts['[Content_Types].xml']=ET.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
settings=ET.fromstring(parts['word/settings.xml'])
ET.SubElement(settings,'{'+W+'}embedTrueTypeFonts')
parts['word/settings.xml']=ET.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
print(OUT)
print('Paragraphs:',len(doc.paragraphs),'Characters:',sum(len(p.text) for p in doc.paragraphs),'Bytes:',OUT.stat().st_size)
